from __future__ import annotations

import io
import hashlib
import html
import logging
import mimetypes
import os
import re
import secrets
import time
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote
from zipfile import BadZipFile, ZipFile

import httpx
import numpy as np
from argon2 import PasswordHasher
from argon2.exceptions import InvalidHashError, VerificationError, VerifyMismatchError
from postgrest.exceptions import APIError
from supabase import Client, create_client
from docx import Document as WordDocument
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image
from pptx import Presentation
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError
from starlette.middleware.sessions import SessionMiddleware
from starlette.responses import Response


BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
load_dotenv(BASE_DIR / ".env")
DATA_DIR = Path(os.getenv("MYDOCX_DATA_DIR", str(BASE_DIR / "data")))
SUPABASE_URL = os.getenv("SUPABASE_URL", "").strip()
SUPABASE_KEY = os.getenv("SUPABASE_KEY", "").strip()
OCR_LANGUAGE = os.getenv("OCR_LANGUAGE", "en").strip() or "en"
if not SUPABASE_URL or not SUPABASE_KEY:
    raise RuntimeError("SUPABASE_URL and SUPABASE_KEY are required.")
supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
logger = logging.getLogger("mydocx.auth")
_ocr_engine: Any | None = None
password_hasher = PasswordHasher()
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
DOCUMENT_BUCKET = "mydocx-documents"
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 180
MAX_RETRIEVAL_DISTANCE = 0.85
MAX_RETRIEVAL_DISTANCE_GAP = 0.04
PROGRESS_TTL_SECONDS = 30 * 60
_progress_updates: dict[str, dict[str, Any]] = {}
SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".csv", ".json", ".html", ".xml",
    ".docx", ".pptx", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp",
}


class ChatRequest(BaseModel):
    question: str = Field(min_length=1, max_length=2000)
    document_ids: list[str] = Field(default_factory=list, max_length=20)
    chat_id: str | None = None
    progress_id: str | None = Field(default=None, max_length=64)


class CreateChatRequest(BaseModel):
    document_ids: list[str] = Field(default_factory=list, max_length=20)


class RegisterRequest(BaseModel):
    display_name: str = Field(min_length=2, max_length=100)
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=12, max_length=128)


class LoginRequest(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=1, max_length=128)


def _update_progress(
    progress_id: str | None,
    user_id: str,
    title: str,
    detail: str,
) -> None:
    if not progress_id:
        return
    now = time.monotonic()
    expired = [
        key for key, value in _progress_updates.items()
        if now - value["updated_at"] > PROGRESS_TTL_SECONDS
    ]
    for key in expired:
        del _progress_updates[key]
    _progress_updates[progress_id] = {
        "user_id": user_id,
        "title": title,
        "detail": detail,
        "updated_at": now,
    }


def _session_secret() -> str:
    configured = os.getenv("SESSION_SECRET_KEY", "").strip()
    if configured:
        if len(configured) < 32:
            raise RuntimeError("SESSION_SECRET_KEY must be at least 32 characters.")
        return configured
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    secret_path = DATA_DIR / ".session_secret"
    try:
        stored = secret_path.read_text(encoding="ascii").strip()
        if len(stored) < 32:
            raise RuntimeError("The stored session secret is invalid. Remove data/.session_secret and restart.")
        return stored
    except FileNotFoundError:
        generated = secrets.token_urlsafe(48)
        try:
            with secret_path.open("x", encoding="ascii") as secret_file:
                secret_file.write(generated)
            return generated
        except FileExistsError:
            stored = secret_path.read_text(encoding="ascii").strip()
            if len(stored) < 32:
                raise RuntimeError("The stored session secret is invalid. Remove data/.session_secret and restart.")
            return stored


def _initialize_database() -> None:
    database_url = os.getenv("SUPABASE_DB_URL", "").strip()
    migration_applied = False

    try:
        for table in ("chat_threads", "chat_documents", "chat_messages"):
            supabase.table(table).select("*").limit(1).execute()
    except Exception as exc:
        if "PGRST205" not in str(exc) or not database_url:
            raise
        _run_database_migration(database_url)
        migration_applied = True

    try:
        supabase.rpc("initialize_mydocx_schema").execute()
    except Exception:
        if not database_url or migration_applied:
            raise
        _run_database_migration(database_url)
        supabase.rpc("initialize_mydocx_schema").execute()


def _run_database_migration(database_url: str) -> None:
    try:
        import psycopg2
        from psycopg2.extensions import ISOLATION_LEVEL_AUTOCOMMIT

        migration = (BASE_DIR / "supabase" / "bootstrap.sql").read_text(encoding="utf-8")
        connection = psycopg2.connect(database_url)
        connection.set_isolation_level(ISOLATION_LEVEL_AUTOCOMMIT)
        try:
            with connection.cursor() as cursor:
                cursor.execute(migration)
        finally:
            connection.close()
    except Exception as exc:
        raise RuntimeError("Automatic database migration failed") from exc


def _database_unavailable(exc: Exception) -> HTTPException:
    message = str(exc)
    logger.error("Supabase database is unavailable", exc_info=(type(exc), exc, exc.__traceback__))
    if "PGRST205" in message:
        detail = (
            "The Supabase chat tables are missing. Run the latest supabase/bootstrap.sql "
            "in the Supabase SQL Editor, then restart the app."
        )
    elif "PGRST202" in message:
        detail = (
            "Supabase schema is not initialized. Run supabase/bootstrap.sql once in the "
            "Supabase SQL Editor, then restart the app."
        )
    elif os.getenv("SUPABASE_DB_URL", "").strip():
        detail = "Database is unavailable. Check SUPABASE_DB_URL and rerun the migration."
    elif SUPABASE_KEY.startswith("sb_publishable_"):
        detail = (
            "SUPABASE_KEY is a publishable browser key. Backend operations require the "
            "service_role key, and the schema must be created first."
        )
    else:
        detail = "Database is unavailable. Check SUPABASE_URL and SUPABASE_KEY."
    return HTTPException(status_code=503, detail=detail)


def get_current_user(request: Request) -> dict[str, Any]:
    user_id = request.session.get("user_id")
    if not isinstance(user_id, str):
        raise HTTPException(status_code=401, detail="Please sign in to continue.")
    try:
        response = (
            supabase.table("users")
            .select("id,display_name,email")
            .eq("id", user_id)
            .eq("is_active", True)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        raise _database_unavailable(exc) from exc
    if not response.data:
        request.session.clear()
        raise HTTPException(status_code=401, detail="Please sign in to continue.")
    return response.data[0]


@asynccontextmanager
async def lifespan(_: FastAPI):
    try:
        _initialize_database()
    except Exception as exc:
        if os.getenv("SUPABASE_DB_URL", "").strip():
            logger.exception("Supabase automatic schema migration failed")
        elif "PGRST205" in str(exc):
            logger.error(
                "Supabase chat tables are missing. In the Supabase Dashboard, open SQL Editor, "
                "run the latest supabase/bootstrap.sql, and restart the app. Set SUPABASE_DB_URL "
                "to enable automatic migration instead."
            )
        else:
            logger.error(
                "Supabase schema is not initialized. In the Supabase Dashboard, open SQL Editor "
                "and run supabase/bootstrap.sql, or set SUPABASE_DB_URL for automatic setup: %s",
                exc,
            )
    yield


app = FastAPI(title="MyDocx.ai Document Chat", lifespan=lifespan)


@app.exception_handler(APIError)
async def handle_supabase_api_error(_: Request, exc: APIError) -> JSONResponse:
    error = _database_unavailable(exc)
    return JSONResponse(status_code=error.status_code, content={"detail": error.detail})


app.add_middleware(
    SessionMiddleware,
    secret_key=_session_secret(),
    session_cookie="mydocx_session",
    max_age=60 * 60 * 24 * 7,
    same_site="lax",
    https_only=os.getenv("MYDOCX_COOKIE_SECURE", "false").lower() == "true",
)
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

def _fallback_embedding(text: str) -> list[float]:
    """Produce a stable offline vector when no embedding provider is configured."""
    vector = [0.0] * 768
    tokens = re.findall(r"[a-z0-9]+", text.casefold())
    for token in tokens:
        digest = uuid.UUID(hashlib.sha256(token.encode()).hexdigest()[:32])
        for offset in range(3):
            index = (digest.int + offset * 7919) % len(vector)
            vector[index] += 1.0 if offset == 0 else 0.35
    norm = sum(value * value for value in vector) ** 0.5
    return [value / norm for value in vector] if norm else vector


async def _embeddings(
    texts: list[str],
    task_type: str = "RETRIEVAL_DOCUMENT",
) -> list[list[float]]:
    if not texts:
        return []
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if api_key:
        try:
            async with httpx.AsyncClient(timeout=30) as client:
                embeddings: list[list[float]] = []
                for offset in range(0, len(texts), 100):
                    batch = texts[offset:offset + 100]
                    response = await client.post(
                        "https://generativelanguage.googleapis.com/v1beta/models/gemini-embedding-001:batchEmbedContents",
                        headers={"x-goog-api-key": api_key},
                        json={
                            "requests": [
                                {
                                    "model": "models/gemini-embedding-001",
                                    "content": {"parts": [{"text": text}]},
                                    "outputDimensionality": 768,
                                    "taskType": task_type,
                                }
                                for text in batch
                            ]
                        },
                    )
                    response.raise_for_status()
                    embeddings.extend(
                        item["values"] for item in response.json()["embeddings"]
                    )
                if len(embeddings) != len(texts):
                    raise ValueError("Gemini returned an unexpected number of embeddings.")
                return embeddings
        except (httpx.HTTPError, KeyError, IndexError, TypeError) as exc:
            logging.getLogger("mydocx.rag").warning(
                "Gemini embedding failed; using offline fallback: %s", exc
            )
        except ValueError as exc:
            logging.getLogger("mydocx.rag").warning(
                "Gemini embedding failed; using offline fallback: %s", exc
            )
    return [_fallback_embedding(text) for text in texts]


async def _embedding(text: str, task_type: str = "RETRIEVAL_DOCUMENT") -> list[float]:
    return (await _embeddings([text], task_type))[0]


def split_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.encode("utf-16-le", "surrogatepass").decode("utf-16-le", "replace")
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(start + size, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(" ", start + size // 2, end)
            if boundary > start:
                end = boundary
        chunks.append(normalized[start:end].strip())
        if end == len(normalized):
            break
        start = max(end - overlap, start + 1)
    return [chunk for chunk in chunks if chunk]


def _page_chunks(text: str) -> list[dict[str, Any]]:
    """Split PDF text by physical page so retrieval can remain page-scoped."""
    matches = list(re.finditer(r"(?m)^\[Page (\d+)\]\s*$", text))
    if not matches:
        return [{"page_number": 1, "text": chunk} for chunk in split_text(text)]

    page_chunks: list[dict[str, Any]] = []
    for index, match in enumerate(matches):
        page_number = int(match.group(1))
        page_end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        page_text = text[match.end():page_end].strip()
        for chunk in split_text(page_text):
            page_chunks.append({"page_number": page_number, "text": chunk})
    return page_chunks


def _document_page_count(
    filename: str,
    data: bytes,
    page_chunks: list[dict[str, Any]],
) -> tuple[int, bool]:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return max(1, len(PdfReader(io.BytesIO(data)).pages)), True
    if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        return 1, True
    if extension == ".pptx":
        return max(1, len(Presentation(io.BytesIO(data)).slides)), True
    return max((item["page_number"] for item in page_chunks), default=1), False


def _append_chunk_without_overlap(previous_text: str, next_chunk: str) -> str:
    overlap = min(len(previous_text), len(next_chunk))
    while overlap and previous_text[-overlap:] != next_chunk[:overlap]:
        overlap -= 1
    return previous_text + next_chunk[overlap:]


def _requests_full_document(question: str) -> bool:
    return bool(re.search(
        r"\b(?:full|entire|complete|verbatim|whole)\b.{0,32}\b(?:text|contents?|output|document|file|pdf|transcript)\b"
        r"|\b(?:show|give|provide|print|return|display|extract|transcribe)\b.{0,32}\b(?:everything|all|full|entire|complete)\b",
        question,
        re.IGNORECASE,
    ))


def _requested_pages(question: str) -> list[int]:
    pages: set[int] = set()
    for match in re.finditer(
        r"\bpages?\s*(\d+)\s*(?:-|–|to)\s*(\d+)", question, re.IGNORECASE
    ):
        start, end = sorted((int(match.group(1)), int(match.group(2))))
        pages.update(range(start, end + 1))
    for match in re.finditer(r"\bpages?\s*(\d+)", question, re.IGNORECASE):
        pages.add(int(match.group(1)))
    for match in re.finditer(r"(?:\band\b|,)\s*(\d+)\b", question, re.IGNORECASE):
        pages.add(int(match.group(1)))
    for match in re.finditer(r"\b(?:on|from)\s+page\s+(\d+)\b", question, re.IGNORECASE):
        pages.add(int(match.group(1)))
    return sorted(pages)


def _requests_document_summary(question: str) -> bool:
    return bool(re.search(
        r"\b(?:summari[sz]e|overview|key points?|main points?|important conclusions?|main ideas?)\b",
        question,
        re.IGNORECASE,
    ))


def _requests_exam_paper(question: str) -> bool:
    return bool(re.search(
        r"\b(?:exam|examination|question)\s+papers?\b"
        r"|\b(?:practice|revision|study|exam|examination|test|quiz)\b.{0,60}\b(?:questions?|quiz(?:zes)?|tests?|papers?)\b"
        r"|\b(?:give|make|create|generate|prepare|write|need|want|ask)\b.{0,40}\b(?:me\s+)?(?:some\s+)?(?:questions?|quiz(?:zes)?|tests?)\b"
        r"|\bquestions?\b.{0,40}\b(?:for\s+)?(?:practice|revision|study|my\s+exam|the\s+exam)\b"
        r"|\b(?:tomorrow|upcoming|next)\b.{0,35}\b(?:exam|test)\b.{0,40}\b(?:questions?|practice)\b",
        question,
        re.IGNORECASE,
    ))


def _requests_answer_key(question: str) -> bool:
    return bool(re.search(
        r"\banswer\s+keys?\b"
        r"|\b(?:answers?|solutions?)\s+(?:to|for)\s+(?:these|those|the|my)\b"
        r"|\b(?:give|show|provide|write|need|want|tell)\b.{0,35}\b(?:answers?|solutions?)\b"
        r"|\bsolve\b.{0,35}\bquestions?\b"
        r"|\banswers?\s+(?:please|now)\b"
        r"|\b(?:answer|answers)\b[.!?]?",
        question,
        re.IGNORECASE,
    ))


def _recent_chat_context(chat_id: str, limit: int = 30) -> str:
    messages = (
        supabase.table("chat_messages")
        .select("id,role,content")
        .eq("chat_id", chat_id)
        .order("created_at", desc=True)
        .order("id", desc=True)
        .limit(limit)
        .execute()
    ).data
    recent = list(reversed(messages))
    return "\n\n".join(
        f"{'User' if message['role'] == 'user' else 'Assistant'}: {message['content'][:6000]}"
        for message in recent
    )[-22000:]


def _ocr_image(image: Image.Image) -> str:
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from paddleocr import PaddleOCR
        except ModuleNotFoundError as exc:
            if exc.name != "paddleocr":
                raise
            raise HTTPException(
                status_code=503,
                detail=(
                    "OCR support is not installed. Activate the project's .venv and run "
                    "python -m pip install -r requirements.txt, then restart the app."
                ),
            ) from exc
        _ocr_engine = PaddleOCR(
            lang=OCR_LANGUAGE,
            device="cpu",
            enable_mkldnn=False,
            text_detection_model_name="PP-OCRv6_small_det",
            text_recognition_model_name="PP-OCRv6_small_rec",
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
        )
    results = _ocr_engine.predict(input=np.asarray(image.convert("RGB")))
    recognized_text: list[str] = []
    for result in results:
        result_data = result.json["res"]
        recognized_text.extend(
            text.strip()
            for text in result_data["rec_texts"]
            if isinstance(text, str) and text.strip()
        )
    return "\n".join(recognized_text)


def _extract_office_images(data: bytes, media_folder: str) -> list[str]:
    extracted: list[str] = []
    with ZipFile(io.BytesIO(data)) as archive:
        image_paths = [
            name for name in archive.namelist()
            if name.startswith(media_folder) and Path(name).suffix.lower()
            in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
        ]
        for image_path in image_paths:
            image_text = _ocr_image(Image.open(io.BytesIO(archive.read(image_path))))
            if image_text:
                extracted.append(f"[Embedded image: {Path(image_path).name}]\n{image_text}")
    return extracted


def _extract_pdf(data: bytes) -> tuple[str, list[str]]:
    reader = PdfReader(io.BytesIO(data))
    pages: list[str] = []
    warnings: list[str] = []
    import pymupdf as fitz

    pdf = fitz.open(stream=data, filetype="pdf")
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        image_texts: list[str] = []
        for image_index, image_info in enumerate(pdf[page_number - 1].get_images(full=True), start=1):
            try:
                image_data = pdf.extract_image(image_info[0])["image"]
                image_text = _ocr_image(Image.open(io.BytesIO(image_data)))
                if image_text:
                    image_texts.append(f"[Image {image_index}]\n{image_text}")
            except HTTPException as exc:
                warnings.append(f"Page {page_number}: {exc.detail}")
        if len(page_text.strip()) < 40 and not image_texts:
            try:
                page_image = pdf[page_number - 1].get_pixmap(matrix=fitz.Matrix(2, 2))
                page_ocr = _ocr_image(Image.open(io.BytesIO(page_image.tobytes("png"))))
                if page_ocr:
                    image_texts.append(f"[Page OCR]\n{page_ocr}")
            except HTTPException as exc:
                warnings.append(f"Page {page_number}: {exc.detail}")
        content = "\n".join(part for part in [page_text, *image_texts] if part.strip())
        pages.append(f"[Page {page_number}]\n{content}".rstrip())
    pdf.close()
    return "\n\n".join(pages), warnings


def extract_document(filename: str, data: bytes) -> tuple[str, list[str]]:
    extension = Path(filename).suffix.lower()
    warnings: list[str] = []
    try:
        if extension == ".pdf":
            return _extract_pdf(data)
        if extension in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
            return _ocr_image(Image.open(io.BytesIO(data))), warnings
        if extension == ".docx":
            document = WordDocument(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for section in document.sections:
                paragraphs.extend(
                    paragraph.text
                    for paragraph in (*section.header.paragraphs, *section.footer.paragraphs)
                    if paragraph.text.strip()
                )
            tables = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            image_texts = _extract_office_images(data, "word/media/")
            return "\n".join(paragraphs + tables + image_texts), warnings
        if extension == ".pptx":
            presentation = Presentation(io.BytesIO(data))
            slides: list[str] = []
            for index, slide in enumerate(presentation.slides, start=1):
                slide_text: list[str] = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        slide_text.append(shape.text)
                    if getattr(shape, "has_table", False):
                        slide_text.extend(
                            " | ".join(cell.text.strip() for cell in row.cells)
                            for row in shape.table.rows
                        )
                slides.append(f"[Slide {index}]\n" + "\n".join(slide_text))
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slides.append(f"[Slide {index} speaker notes]\n{notes}")
            image_texts = _extract_office_images(data, "ppt/media/")
            return "\n\n".join(slides + image_texts), warnings
        if extension in {".txt", ".md", ".csv", ".json", ".html", ".xml"}:
            text = data.decode("utf-8-sig")
            if extension == ".html":
                text = html.unescape(re.sub(r"<[^>]+>", " ", text))
            return text, warnings
    except HTTPException:
        raise
    except (
        BadZipFile,
        FileNotDecryptedError,
        OSError,
        UnicodeDecodeError,
        ValueError,
        KeyError,
        IndexError,
        PdfReadError,
    ) as exc:
        raise HTTPException(status_code=422, detail=f"Could not read {filename}: {exc}") from exc
    raise HTTPException(status_code=415, detail=f"Unsupported file type: {extension or 'unknown'}")


@app.get("/")
async def home(request: Request) -> Response:
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/app")
async def workspace(request: Request) -> Response:
    if not isinstance(request.session.get("user_id"), str):
        return RedirectResponse("/login", status_code=303)
    return FileResponse(STATIC_DIR / "chat.html")


@app.get("/login")
async def login_page(request: Request) -> Response:
    if isinstance(request.session.get("user_id"), str):
        return RedirectResponse("/app", status_code=303)
    return FileResponse(STATIC_DIR / "login.html")


@app.get("/register")
async def register_page(request: Request) -> Response:
    if isinstance(request.session.get("user_id"), str):
        return RedirectResponse("/app", status_code=303)
    return FileResponse(STATIC_DIR / "register.html")


@app.post("/api/auth/register")
async def register(body: RegisterRequest, request: Request) -> dict[str, Any]:
    display_name = " ".join(body.display_name.split())
    email = body.email.strip().lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise HTTPException(status_code=422, detail="Enter a valid email address.")
    if len(display_name) < 2:
        raise HTTPException(status_code=422, detail="Name must contain at least 2 characters.")

    password_hash = password_hasher.hash(body.password)
    try:
        response = supabase.table("users").insert({
            "display_name": display_name,
            "email": email,
            "password_hash": password_hash,
        }).execute()
    except Exception as exc:
        if "23505" in str(exc):
            raise HTTPException(status_code=409, detail="An account with that email already exists.") from exc
        raise _database_unavailable(exc) from exc
    user_id = response.data[0]["id"]

    request.session.clear()
    request.session["user_id"] = str(user_id)
    return {"user": {"id": str(user_id), "display_name": display_name, "email": email}}


@app.post("/api/auth/login")
async def login(body: LoginRequest, request: Request) -> dict[str, Any]:
    email = body.email.strip().lower()
    try:
        response = (
            supabase.table("users")
            .select("id,display_name,email,password_hash")
            .eq("email", email)
            .eq("is_active", True)
            .limit(1)
            .execute()
        )
    except Exception as exc:
        raise _database_unavailable(exc) from exc

    if not response.data:
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    user = response.data[0]
    try:
        password_hasher.verify(user["password_hash"], body.password)
    except (VerifyMismatchError, VerificationError, InvalidHashError) as exc:
        raise HTTPException(status_code=401, detail="Email or password is incorrect.") from exc

    request.session.clear()
    request.session["user_id"] = str(user["id"])
    return {
        "user": {
            "id": str(user["id"]),
            "display_name": user["display_name"],
            "email": user["email"],
        }
    }


@app.post("/api/auth/logout")
async def logout(request: Request) -> dict[str, str]:
    request.session.clear()
    return {"status": "signed out"}


@app.get("/api/auth/me")
async def current_user(user: dict[str, Any] = Depends(get_current_user)) -> dict[str, Any]:
    return {"user": user}


@app.get("/api/health")
async def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/api/progress/{progress_id}")
async def get_progress(
    progress_id: str,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, str]:
    progress = _progress_updates.get(progress_id)
    if progress is None:
        return {"title": "Starting", "detail": "Preparing your request."}
    if time.monotonic() - progress["updated_at"] > PROGRESS_TTL_SECONDS:
        del _progress_updates[progress_id]
        return {"title": "Starting", "detail": "Preparing your request."}
    if progress["user_id"] != user["id"]:
        return {"title": "Starting", "detail": "Preparing your request."}
    return {"title": progress["title"], "detail": progress["detail"]}


@app.delete("/api/progress/{progress_id}")
async def clear_progress(
    progress_id: str,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, str]:
    progress = _progress_updates.get(progress_id)
    if progress and progress["user_id"] == user["id"]:
        del _progress_updates[progress_id]
    return {"status": "cleared"}


@app.get("/api/documents")
async def list_documents(
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, list[dict[str, Any]]]:
    response = (
        supabase.rpc("list_mydocx_documents", {"p_user_id": user["id"]})
        .execute()
    )
    return {"documents": response.data}


def _owned_document(document_id: str, user_id: str) -> dict[str, Any]:
    response = (
        supabase.table("documents")
        .select("id,filename,page_count,page_count_known,storage_path")
        .eq("id", document_id)
        .eq("user_id", user_id)
        .limit(1)
        .execute()
    )
    if not response.data:
        raise HTTPException(status_code=404, detail="Document not found.")
    return response.data[0]


@app.get("/api/documents/{document_id}/content")
async def get_document_content(
    document_id: str,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    document = _owned_document(document_id, user["id"])
    records = supabase.rpc("get_mydocx_chunks", {
        "p_user_id": user["id"],
        "p_document_ids": [document_id],
        "p_page_numbers": None,
    }).execute().data
    pages: dict[int, list[tuple[int, str]]] = {}
    for record in records:
        pages.setdefault(record["page_number"], []).append(
            (record["chunk_index"], record["content"])
        )
    content = "\n\n".join(
        f"[Page {page_number}]\n" + _combine_chunks(chunks)
        for page_number, chunks in sorted(pages.items())
    )
    return {
        "filename": document["filename"],
        "page_count": document["page_count"],
        "page_count_known": document["page_count_known"],
        "text": content,
    }


@app.get("/api/documents/{document_id}/file")
async def get_document_file(
    document_id: str,
    download: bool = False,
    user: dict[str, Any] = Depends(get_current_user),
) -> Response:
    document = _owned_document(document_id, user["id"])
    storage_path = document["storage_path"]
    if not storage_path:
        raise HTTPException(
            status_code=404,
            detail="The original file is not available. Re-upload this document to open it.",
        )
    file_data = supabase.storage.from_(DOCUMENT_BUCKET).download(storage_path)
    filename = quote(document["filename"])
    disposition = "attachment" if download else "inline"
    return Response(
        content=file_data,
        media_type=mimetypes.guess_type(document["filename"])[0] or "application/octet-stream",
        headers={
            "Content-Disposition": f"{disposition}; filename*=UTF-8''{filename}",
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


def _combine_chunks(chunks: list[tuple[int, str]]) -> str:
    content = ""
    for _, chunk_text in sorted(chunks):
        content = _append_chunk_without_overlap(content, chunk_text)
    return content


def _owned_chat(chat_id: str, user_id: str) -> dict[str, Any]:
    response = (
        supabase.table("chat_threads")
        .select("id,title,description")
        .eq("id", chat_id)
        .eq("user_id", user_id)
        .limit(1)
        .execute()
    )
    if not response.data:
        raise HTTPException(status_code=404, detail="Chat not found.")
    return response.data[0]


def _attach_chat_documents(chat_id: str, user_id: str, document_ids: list[str]) -> list[dict[str, Any]]:
    unique_ids = list(dict.fromkeys(document_ids))
    if not unique_ids:
        return []
    owned = (
        supabase.table("documents")
        .select("id,filename")
        .eq("user_id", user_id)
        .in_("id", unique_ids)
        .execute()
    ).data
    documents_by_id = {document["id"]: document for document in owned}
    if len(documents_by_id) != len(unique_ids):
        raise HTTPException(status_code=404, detail="One or more documents are not available.")

    linked = (
        supabase.table("chat_documents")
        .select("document_id")
        .eq("chat_id", chat_id)
        .in_("document_id", unique_ids)
        .execute()
    ).data
    linked_ids = {row["document_id"] for row in linked}
    new_links = [
        {"chat_id": chat_id, "document_id": document_id}
        for document_id in unique_ids
        if document_id not in linked_ids
    ]
    if new_links:
        supabase.table("chat_documents").insert(new_links).execute()

    all_links = (
        supabase.table("chat_documents")
        .select("document_id")
        .eq("chat_id", chat_id)
        .execute()
    ).data
    all_document_ids = [row["document_id"] for row in all_links]
    all_documents = (
        supabase.table("documents")
        .select("id,filename")
        .eq("user_id", user_id)
        .in_("id", all_document_ids)
        .execute()
    ).data if all_document_ids else []
    ordered_documents = sorted(all_documents, key=lambda document: document["filename"].casefold())
    title = ", ".join(document["filename"] for document in ordered_documents)
    supabase.table("chat_threads").update({
        "title": title[:200] or "New chat",
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }).eq("id", chat_id).eq("user_id", user_id).execute()
    return [documents_by_id[document_id] for document_id in unique_ids]


@app.post("/api/chats")
async def create_chat(
    body: CreateChatRequest,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    chat_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    supabase.table("chat_threads").insert({
        "id": chat_id,
        "user_id": user["id"],
        "title": "New chat",
        "description": "Start a conversation",
        "created_at": now,
        "updated_at": now,
    }).execute()
    documents = _attach_chat_documents(chat_id, user["id"], body.document_ids)
    chat = _owned_chat(chat_id, user["id"])
    return {"chat": {**chat, "documents": documents}}


@app.get("/api/chats")
async def list_chats(
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, list[dict[str, Any]]]:
    chats = (
        supabase.table("chat_threads")
        .select("id,title,description,updated_at")
        .eq("user_id", user["id"])
        .order("updated_at", desc=True)
        .execute()
    ).data
    if not chats:
        return {"chats": []}

    links = (
        supabase.table("chat_documents")
        .select("chat_id,document_id")
        .in_("chat_id", [chat["id"] for chat in chats])
        .execute()
    ).data
    document_ids = list(dict.fromkeys(link["document_id"] for link in links))
    documents = (
        supabase.table("documents")
        .select("id,filename")
        .eq("user_id", user["id"])
        .in_("id", document_ids)
        .execute()
    ).data if document_ids else []
    filenames = {document["id"]: document["filename"] for document in documents}
    documents_by_chat: dict[str, list[str]] = {}
    for link in links:
        filename = filenames.get(link["document_id"])
        if filename:
            documents_by_chat.setdefault(link["chat_id"], []).append(filename)
    for chat in chats:
        chat["documents"] = list(dict.fromkeys(documents_by_chat.get(chat["id"], [])))
    return {"chats": chats}


@app.get("/api/chats/{chat_id}")
async def get_chat(
    chat_id: str,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    chat = _owned_chat(chat_id, user["id"])
    links = (
        supabase.table("chat_documents")
        .select("document_id")
        .eq("chat_id", chat_id)
        .execute()
    ).data
    document_ids = [row["document_id"] for row in links]
    documents = (
        supabase.table("documents")
        .select("id,filename")
        .eq("user_id", user["id"])
        .in_("id", document_ids)
        .execute()
    ).data if document_ids else []
    messages = (
        supabase.table("chat_messages")
        .select("id,role,content,sources,created_at")
        .eq("chat_id", chat_id)
        .order("created_at")
        .order("id")
        .execute()
    ).data
    chat["documents"] = documents
    chat["messages"] = messages
    return {"chat": chat}


@app.delete("/api/chats/{chat_id}")
async def delete_chat(
    chat_id: str,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, str]:
    _owned_chat(chat_id, user["id"])
    supabase.table("chat_threads").delete().eq("id", chat_id).eq("user_id", user["id"]).execute()
    return {"status": "deleted"}


@app.post("/api/chats/{chat_id}/documents")
async def add_chat_documents(
    chat_id: str,
    body: CreateChatRequest,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    _owned_chat(chat_id, user["id"])
    _attach_chat_documents(chat_id, user["id"], body.document_ids)
    return {"status": "attached"}


@app.post("/api/documents")
async def upload_documents(
    files: list[UploadFile] = File(...),
    progress_id: str | None = Header(default=None, alias="X-Progress-ID", max_length=64),
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    results: list[dict[str, Any]] = []
    total_files = len(files)
    for file_index, upload in enumerate(files, start=1):
        filename = Path(upload.filename or "document").name
        extension = Path(filename).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise HTTPException(status_code=415, detail=f"Unsupported file type: {extension or 'unknown'}")
        _update_progress(
            progress_id,
            user["id"],
            "Reading your document",
            f"Checking file {file_index} of {total_files}: {filename}.",
        )
        data = await upload.read(MAX_UPLOAD_BYTES + 1)
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail=f"{filename} exceeds the 20 MB upload limit.")
        _update_progress(
            progress_id,
            user["id"],
            "Extracting document text",
            f"Reading {filename}; scanned pages and images are being processed with OCR.",
        )
        text, warnings = extract_document(filename, data)
        _update_progress(
            progress_id,
            user["id"],
            "Preparing searchable passages",
            f"Splitting the text from {filename} into passages.",
        )
        page_chunks = _page_chunks(text)
        chunks = [item["text"] for item in page_chunks]
        if not chunks:
            raise HTTPException(status_code=422, detail=f"No readable text found in {filename}.")
        document_id = str(uuid.uuid4())
        page_count, page_count_known = _document_page_count(filename, data, page_chunks)
        storage_path = f"{user['id']}/{document_id}{extension}"
        content_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"
        _update_progress(
            progress_id,
            user["id"],
            "Saving the original file",
            f"Storing {filename} securely so it can be opened later.",
        )
        supabase.storage.from_(DOCUMENT_BUCKET).upload(
            storage_path,
            data,
            {"content-type": content_type, "upsert": "false"},
        )
        document_inserted = False
        try:
            supabase.table("documents").insert({
                "id": document_id,
                "user_id": user["id"],
                "filename": filename,
                "page_count": page_count,
                "page_count_known": page_count_known,
                "storage_path": storage_path,
            }).execute()
            document_inserted = True
            _update_progress(
                progress_id,
                user["id"],
                "Building the document search index",
                f"Creating embeddings for {len(chunks)} passages from {filename}.",
            )
            embeddings = await _embeddings(chunks)
            chunk_rows = [
                {
                    "document_id": document_id,
                    "chunk_index": index,
                    "page_number": page_chunks[index]["page_number"],
                    "content": chunk,
                    "embedding": embedding,
                }
                for index, (chunk, embedding) in enumerate(zip(chunks, embeddings))
            ]
            supabase.table("document_chunks").insert(chunk_rows).execute()
        except Exception:
            logger.exception("Failed to index uploaded document %s", filename)
            try:
                if document_inserted:
                    supabase.table("documents").delete().eq("id", document_id).eq("user_id", user["id"]).execute()
                supabase.storage.from_(DOCUMENT_BUCKET).remove([storage_path])
            except Exception:
                logger.exception("Failed to clean up uploaded document %s", filename)
            raise
        results.append({
            "id": document_id,
            "filename": filename,
            "chunks": len(chunks),
            "pages": page_count if page_count_known else None,
            "warnings": warnings,
        })
    _update_progress(
        progress_id,
        user["id"],
        "Document ready",
        f"Finished processing {total_files} document{'s' if total_files != 1 else ''}.",
    )
    return {"documents": results}


@app.delete("/api/documents/{document_id}")
async def delete_document(
    document_id: str,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, str]:
    document = _owned_document(document_id, user["id"])
    if document["storage_path"]:
        supabase.storage.from_(DOCUMENT_BUCKET).remove([document["storage_path"]])
    supabase.table("documents").delete().eq("id", document_id).eq("user_id", user["id"]).execute()
    return {"status": "deleted"}


async def _generate_gemini_answer(
    question: str,
    passages: list[dict[str, Any]],
    api_key: str,
    conversation_context: str = "",
) -> str:
    primary_model = os.getenv("GEMINI_MODEL", "gemini-flash-lite-latest").strip() or "gemini-flash-lite-latest"
    fallback_model = (
        os.getenv("GEMINI_FALLBACK_MODEL", "gemini-flash-latest").strip()
        or "gemini-flash-latest"
    )
    models = list(dict.fromkeys((primary_model, fallback_model)))
    context = "\n\n".join(
        f"Source: {item['filename']} · {item.get('page_or_slide') or 'Page not recorded'}\n"
        f"{item['text'][:12000]}"
        for item in passages
    ) or "No relevant document passages were found."
    if _requests_answer_key(question):
        instructions = (
            "The user wants answers to practice or exam questions. Use the recent conversation to "
            "identify the latest set of questions, then provide a numbered answer key in the same "
            "order. Explain each answer briefly and base factual claims only on the document excerpts. "
            "If the recent conversation does not contain a question set, create a short practice set "
            "from the notes and provide its answers. Do not claim the notes contain existing exam questions."
        )
    elif _requests_exam_paper(question):
        instructions = (
            "Create useful questions for exam practice based on the supplied notes, even if the notes "
            "do not contain pre-written exam questions. Create 10 questions grounded in the material: "
            "include a mix of short-answer, explain/compare, and application questions where suitable. "
            "Order them from easier to harder, number them, and assign marks for a total of 30. "
            "Do not include answers unless requested. Never say the notes contain no questions; "
            "write new questions from the concepts and facts in the notes without inventing unsupported facts."
        )
    elif _requests_document_summary(question):
        instructions = (
            "Summarize the supplied document in plain, simple English. Rewrite difficult sentences "
            "so they are easy to understand. Start with one short overview, then give 5-8 concise "
            "bullets covering the most important distinct ideas. Keep names, key terms, and facts "
            "accurate. Do not copy long passages, repeat ideas, or include unrelated information. "
            "If the document is too long, prioritize its central themes and important supporting points."
        )
    else:
        instructions = (
            "Answer accurately and concisely using only the supplied document excerpts. Treat the "
            "source page labels as strict scope: never use facts from another page or document. "
            "Never infer facts that are not in the excerpts. If the excerpts do not contain the answer, "
            "say so. Give a simple explanation, then cite the source filename and page."
        )
    prompt = (
        f"You are MyDocx.ai, a careful document-grounded assistant. {instructions}\n\n"
        f"Conversation memory from this chat, including earlier questions and assistant replies "
        f"(use it to understand follow-ups and references; document excerpts remain the source "
        f"of truth for document facts):\n"
        f"{conversation_context or 'No earlier chat messages.'}\n\n"
        f"Document excerpts:\n{context}\n\nQuestion: {question}"
    )
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            interaction: dict[str, Any] | None = None
            for index, model in enumerate(models):
                response = await client.post(
                    "https://generativelanguage.googleapis.com/v1beta/interactions",
                    headers={"x-goog-api-key": api_key},
                    json={
                        "model": model,
                        "input": prompt,
                        "store": False,
                        "generation_config": {
                            "temperature": 0.2,
                            "max_output_tokens": (
                                1600
                                if _requests_exam_paper(question) or _requests_answer_key(question)
                                else 450
                            ),
                        },
                    },
                )
                try:
                    response.raise_for_status()
                except httpx.HTTPStatusError as exc:
                    retryable_statuses = {400, 404, 408, 429, 500, 502, 503, 504}
                    if index + 1 < len(models) and exc.response.status_code in retryable_statuses:
                        logging.getLogger("mydocx.ai").warning(
                            "Gemini model %s returned HTTP %s; retrying with %s",
                            model,
                            exc.response.status_code,
                            models[index + 1],
                        )
                        continue
                    raise
                interaction = response.json()
                break
            if interaction is None:
                raise HTTPException(status_code=502, detail="Gemini returned no answer for this request.")
            output_steps = [
                step for step in interaction.get("steps", [])
                if step.get("type") == "model_output"
            ]
            if not output_steps:
                raise HTTPException(status_code=502, detail="Gemini returned no answer for this request.")
            answer = "\n".join(
                part["text"]
                for step in output_steps
                for part in step.get("content", [])
                if isinstance(part.get("text"), str)
            ).strip()
            if not answer:
                raise HTTPException(status_code=502, detail="Gemini returned an empty answer.")
            return answer
    except httpx.HTTPStatusError as exc:
        detail = "Gemini API request failed."
        try:
            detail = exc.response.json().get("error", {}).get("message", detail)
        except ValueError:
            pass
        raise HTTPException(status_code=502, detail=detail) from exc
    except httpx.RequestError as exc:
        raise HTTPException(status_code=502, detail=f"Could not connect to Gemini: {exc}") from exc
    except (KeyError, IndexError, TypeError, ValueError, AttributeError) as exc:
        raise HTTPException(status_code=502, detail=f"Gemini returned an unexpected response: {exc}") from exc


async def _generate_answer(
    question: str,
    passages: list[dict[str, Any]],
    conversation_context: str = "",
) -> str:
    gemini_api_key = os.getenv("GEMINI_API_KEY")
    if gemini_api_key:
        try:
            return await _generate_gemini_answer(
                question,
                passages,
                gemini_api_key,
                conversation_context,
            )
        except HTTPException as exc:
            if exc.status_code not in {502, 503, 504}:
                raise
            logging.getLogger("mydocx.ai").warning(
                "Gemini was unavailable while generating a document-grounded response."
            )
            if _requests_exam_paper(question) or _requests_answer_key(question):
                raise

            def format_excerpt(passage: dict[str, Any]) -> str:
                source = passage["filename"]
                if passage.get("page_or_slide"):
                    source += f" · {passage['page_or_slide']}"
                return f"**{source}**\n{passage['text']}"

            excerpts = "\n\n".join(
                format_excerpt(passage) for passage in passages
            )
            if excerpts:
                return (
                    "The AI service is temporarily busy, so here are the relevant passages "
                    "I found in your documents:\n\n"
                    f"{excerpts}"
                )
            return "The AI service is temporarily unavailable, and I couldn't find a relevant passage."

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        if _requests_exam_paper(question) or _requests_answer_key(question):
            return (
                "Practice questions and answer keys need a configured AI provider. "
                "Add GEMINI_API_KEY to your .env file, restart the app, and ask again."
            )
        if not passages:
            return "I couldn't find a relevant passage. Try rephrasing your question or upload a document that covers the topic."
        count = len(passages)
        return (
            f"I found {count} relevant passage{'s' if count != 1 else ''}. "
            "Add GEMINI_API_KEY to your .env file for a concise, AI-generated answer. "
            "Expand Sources below to read the document excerpts."
        )

    base_url = os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
    model = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
    context = "\n\n".join(
        f"Source: {item['filename']}\n{item['text']}" for item in passages
    ) or "No relevant document passages were found."
    try:
        async with httpx.AsyncClient(timeout=60) as client:
            response = await client.post(
                f"{base_url}/chat/completions",
                headers={"Authorization": f"Bearer {api_key}"},
                json={
                    "model": model,
                    "messages": [
                        {
                            "role": "system",
                            "content": (
                                "Answer using only supplied document excerpts. If asked to create "
                                "practice questions, create 10 useful questions from the notes, ordered "
                                "from easier to harder, with a mix of short-answer and application types "
                                "and marks totaling 30; do not include answers unless asked. If asked "
                                "for answers, use recent chat context for the latest question set and "
                                "provide a numbered key with brief explanations. Do not invent unsupported "
                                "facts. Cite source filenames in your answer."
                            ),
                        },
                        {
                            "role": "user",
                            "content": (
                                f"Conversation memory from this chat:\n"
                                f"{conversation_context or 'No earlier chat messages.'}"
                                f"\n\nDocument excerpts:\n{context}\n\nQuestion: {question}"
                            ),
                        },
                    ],
                    "temperature": 0.2,
                },
            )
            response.raise_for_status()
            return response.json()["choices"][0]["message"]["content"]
    except (httpx.HTTPError, KeyError, IndexError, TypeError) as exc:
        raise HTTPException(status_code=502, detail=f"Language model request failed: {exc}") from exc


def _save_chat_message(
    chat_id: str,
    role: str,
    content: str,
    sources: list[dict[str, Any]] | None = None,
) -> None:
    supabase.table("chat_messages").insert({
        "chat_id": chat_id,
        "role": role,
        "content": content,
        "sources": sources or [],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }).execute()
    supabase.table("chat_threads").update({
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }).eq("id", chat_id).execute()


def _chat_reply(chat_id: str, answer: str, sources: list[dict[str, Any]]) -> dict[str, Any]:
    _save_chat_message(chat_id, "assistant", answer, sources)
    document_links = (
        supabase.table("chat_documents")
        .select("document_id")
        .eq("chat_id", chat_id)
        .execute()
    ).data
    return {
        "chat_id": chat_id,
        "document_ids": [link["document_id"] for link in document_links],
        "answer": answer,
        "sources": sources,
    }


@app.post("/api/chat")
async def chat(
    request: ChatRequest,
    user: dict[str, Any] = Depends(get_current_user),
) -> dict[str, Any]:
    _update_progress(
        request.progress_id,
        user["id"],
        "Loading this chat",
        "Restoring its linked documents and conversation context.",
    )
    if request.chat_id:
        thread = _owned_chat(request.chat_id, user["id"])
        _attach_chat_documents(request.chat_id, user["id"], request.document_ids)
        chat_id = request.chat_id
        linked_documents = (
            supabase.table("chat_documents")
            .select("document_id")
            .eq("chat_id", chat_id)
            .execute()
        ).data
        document_ids = [row["document_id"] for row in linked_documents]
        if not document_ids:
            documents = (
                supabase.table("documents")
                .select("id")
                .eq("user_id", user["id"])
                .execute()
            ).data
            document_ids = [document["id"] for document in documents]
            if document_ids:
                _attach_chat_documents(chat_id, user["id"], document_ids)
    else:
        document_ids = request.document_ids
        if not document_ids:
            documents = (
                supabase.table("documents")
                .select("id")
                .eq("user_id", user["id"])
                .execute()
            ).data
            document_ids = [document["id"] for document in documents]
        if not document_ids:
            raise HTTPException(status_code=400, detail="Upload or select a document before chatting.")
        chat_id = str(uuid.uuid4())
        now = datetime.now(timezone.utc).isoformat()
        supabase.table("chat_threads").insert({
            "id": chat_id,
            "user_id": user["id"],
            "title": "New chat",
            "description": "Start a conversation",
            "created_at": now,
            "updated_at": now,
        }).execute()
        _attach_chat_documents(chat_id, user["id"], document_ids)
        thread = _owned_chat(chat_id, user["id"])

    specific_chat_context = bool(request.chat_id or request.document_ids)
    query_document_ids = document_ids or None
    requested_pages = _requested_pages(request.question) or None
    _update_progress(
        request.progress_id,
        user["id"],
        "Checking saved document passages",
        "Confirming that this chat has searchable document content.",
    )
    available = supabase.rpc("count_mydocx_chunks", {
        "p_user_id": user["id"],
        "p_document_ids": query_document_ids,
        "p_page_numbers": requested_pages,
    }).execute()
    chunk_count = int(available.data or 0)
    if chunk_count == 0:
        if document_ids:
            raise HTTPException(status_code=404, detail="The documents in this chat have no searchable passages.")
        raise HTTPException(status_code=400, detail="Upload at least one document before chatting.")
    conversation_context = _recent_chat_context(chat_id)
    if thread["description"] == "Start a conversation":
        supabase.table("chat_threads").update({
            "description": request.question[:500],
        }).eq("id", chat_id).eq("user_id", user["id"]).execute()
    _save_chat_message(chat_id, "user", request.question)

    if document_ids and (
        _requests_document_summary(request.question)
        or _requests_exam_paper(request.question)
        or _requests_answer_key(request.question)
        or requested_pages
    ):
        _update_progress(
            request.progress_id,
            user["id"],
            "Reading the saved notes",
            "Loading the complete note text for your summary, practice questions, or answers.",
        )
        records = supabase.rpc("get_mydocx_chunks", {
            "p_user_id": user["id"],
            "p_document_ids": query_document_ids,
            "p_page_numbers": requested_pages,
        }).execute()
        documents: dict[str, dict[str, Any]] = {}
        for record in records.data:
            text = record["content"]
            metadata = record
            document = documents.setdefault(metadata["document_id"], {
                "filename": metadata["filename"],
                "chunks": [],
            })
            document["chunks"].append((metadata["chunk_index"], text))
        summary_passages: list[dict[str, Any]] = []
        for document_id, document in documents.items():
            content = ""
            for _, chunk_text in sorted(document["chunks"], key=lambda item: item[0]):
                content = _append_chunk_without_overlap(content, chunk_text)
            if content.strip():
                summary_passages.append({
                    "filename": document["filename"],
                    "text": content.strip(),
                    "document_id": document_id,
                    "page_or_slide": f"Page {requested_pages[0]}" if requested_pages and len(requested_pages) == 1 else None,
                    "relevance": 1.0,
                })
        if not summary_passages:
            return _chat_reply(chat_id, "I couldn't extract readable text from the attached document.", [])
        _update_progress(
            request.progress_id,
            user["id"],
            "Preparing your response",
            "Using the document and recent chat messages to compose a grounded response.",
        )
        answer = await _generate_answer(
            request.question,
            summary_passages,
            conversation_context,
        )
        return _chat_reply(chat_id, answer, summary_passages)
    if _requests_full_document(request.question):
        _update_progress(
            request.progress_id,
            user["id"],
            "Reading the requested document",
            "Combining the saved passages in their original order.",
        )
        complete_records = supabase.rpc("get_mydocx_chunks", {
            "p_user_id": user["id"],
            "p_document_ids": query_document_ids,
            "p_page_numbers": requested_pages,
        }).execute()
        grouped_documents: dict[str, dict[str, Any]] = {}
        for record in complete_records.data:
            text = record["content"]
            metadata = record
            document = grouped_documents.setdefault(metadata["document_id"], {
                "filename": metadata["filename"],
                "chunks": [],
            })
            document["chunks"].append((metadata["chunk_index"], text))
        unique_documents: list[dict[str, str]] = []
        seen_documents: set[tuple[str, str]] = set()
        for document in grouped_documents.values():
            content = ""
            for _, chunk_text in sorted(document["chunks"], key=lambda item: item[0]):
                content = _append_chunk_without_overlap(content, chunk_text)
            identity = (document["filename"], content)
            if content.strip() and identity not in seen_documents:
                seen_documents.add(identity)
                unique_documents.append({
                    "filename": document["filename"],
                    "text": content.strip(),
                })
        if not unique_documents:
            return _chat_reply(chat_id, "I couldn't extract readable text from your documents.", [])
        question_lower = request.question.casefold()
        if re.search(r"\ball\s+(?:of\s+)?(?:documents?|files?)\b", question_lower):
            selected_documents = unique_documents
        else:
            selected_documents = [
                document for document in unique_documents
                if document["filename"].casefold() in question_lower
            ]
            if not selected_documents and len(unique_documents) == 1:
                selected_documents = unique_documents
        if not selected_documents:
            filenames = ", ".join(document["filename"] for document in unique_documents)
            return _chat_reply(chat_id, f"I found multiple documents. Which one should I show in full? {filenames}", [])
        answer = "\n\n".join(
            f"## Complete extracted text: {document['filename']}\n\n{document['text']}"
            for document in selected_documents
        )
        return _chat_reply(chat_id, answer, [])
    _update_progress(
        request.progress_id,
        user["id"],
        "Finding relevant passages",
        "Searching the documents linked to this chat.",
    )
    query_embedding = await _embedding(request.question, "RETRIEVAL_QUERY")
    results = supabase.rpc("search_mydocx_chunks", {
        "p_user_id": user["id"],
        "p_document_ids": query_document_ids,
        "p_page_numbers": requested_pages,
        "p_embedding": query_embedding,
        "p_match_count": min(8, chunk_count),
    }).execute()
    retrieved: list[dict[str, Any]] = []
    for record in results.data:
        distance = float(record["distance"])
        retrieved.append({
            "text": record["content"],
            "filename": record["filename"],
            "document_id": record["document_id"],
            "chunk": record["chunk_index"],
            "page_or_slide": f"Page {record['page_number']}",
            "distance": distance,
            "relevance": round(max(0.0, 1.0 - distance), 3),
        })
    if not retrieved:
        return _chat_reply(chat_id, "I couldn't find a relevant passage for that question in your documents. Try rephrasing it or naming the document.", [])
    best_distance = min(passage["distance"] for passage in retrieved)
    max_distance = best_distance + MAX_RETRIEVAL_DISTANCE_GAP
    if not specific_chat_context:
        max_distance = min(MAX_RETRIEVAL_DISTANCE, max_distance)
    relevant: list[dict[str, Any]] = []
    seen_chunks: set[tuple[str, int, str]] = set()
    for passage in retrieved:
        identity = (passage["filename"], passage["chunk"], passage["text"])
        if passage["distance"] <= max_distance and identity not in seen_chunks:
            seen_chunks.add(identity)
            relevant.append(passage)
    if not relevant:
        return _chat_reply(chat_id, "I couldn't find a relevant passage for that question in your documents. Try rephrasing it or naming the document.", [])
    retrieved = relevant
    retrieved.sort(key=lambda passage: (passage["document_id"], passage["chunk"]))
    passages: list[dict[str, Any]] = []
    for passage in retrieved:
        if (
            passages
            and passages[-1]["document_id"] == passage["document_id"]
            and passage["chunk"] == passages[-1]["last_chunk"] + 1
        ):
            passages[-1]["text"] = _append_chunk_without_overlap(
                passages[-1]["text"],
                passage["text"],
            )
            passages[-1]["last_chunk"] = passage["chunk"]
            passages[-1]["relevance"] = max(passages[-1]["relevance"], passage["relevance"])
        else:
            passages.append({
                **passage,
                "first_chunk": passage["chunk"],
                "last_chunk": passage["chunk"],
            })
    passages.sort(key=lambda passage: passage["relevance"], reverse=True)
    passages = passages[:4]
    for passage in passages:
        page_numbers = list(dict.fromkeys(re.findall(r"\[Page (\d+)\]", passage["text"])))
        if len(page_numbers) == 1:
            passage["page_or_slide"] = f"Page {page_numbers[0]}"
        elif page_numbers:
            passage["page_or_slide"] = f"Pages {', '.join(page_numbers)}"
        passage["text"] = re.sub(r"\[Page \d+\]\s*", "", passage["text"])
        passage["text"] = passage["text"][:1800].strip()
        passage.pop("distance")
        passage.pop("document_id")
        passage.pop("chunk")
        passage.pop("first_chunk")
        passage.pop("last_chunk")
    _update_progress(
        request.progress_id,
        user["id"],
        "Writing your answer",
        "Preparing a clear response from the passages found in your documents.",
    )
    answer = await _generate_answer(
        request.question,
        passages,
        conversation_context,
    )
    return _chat_reply(chat_id, answer, passages)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
